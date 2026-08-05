import csv
import json
from collections.abc import Callable
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any

import fitz
from docx import Document
from openpyxl import load_workbook

from researchos.domain.enums import FormatKind


def extract_deterministic_metadata(
    kind: FormatKind,
    path: Path,
    original_filename: str,
) -> dict[str, Any]:
    base: dict[str, Any] = {
        "original_filename": original_filename,
        "extension": Path(original_filename).suffix.lower(),
    }
    extractor: dict[FormatKind, Callable[[Path], dict[str, Any]]] = {
        FormatKind.MARKDOWN: _text_metadata,
        FormatKind.CSV: _csv_metadata,
        FormatKind.DOCX: _docx_metadata,
        FormatKind.XLSX: _xlsx_metadata,
        FormatKind.PDF_TEXT: _pdf_metadata,
    }
    if kind in extractor:
        base.update(extractor[kind](path))
    return _json_safe(base)


def _text_metadata(path: Path) -> dict[str, Any]:
    text = path.read_text(encoding="utf-8-sig")
    return {
        "encoding": "utf-8",
        "line_count": len(text.splitlines()),
        "character_count": len(text),
    }


def _csv_metadata(path: Path) -> dict[str, Any]:
    text = path.read_text(encoding="utf-8-sig")
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample)
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","

    rows = list(csv.reader(text.splitlines(), delimiter=delimiter))
    return {
        "encoding": "utf-8",
        "delimiter": delimiter,
        "row_count": len(rows),
        "column_count": max((len(row) for row in rows), default=0),
        "header": rows[0] if rows else [],
    }


def _docx_metadata(path: Path) -> dict[str, Any]:
    document = Document(path)
    properties = document.core_properties
    return {
        "author": properties.author,
        "title": properties.title,
        "subject": properties.subject,
        "created": properties.created,
        "modified": properties.modified,
        "paragraph_count": len(document.paragraphs),
        "table_count": len(document.tables),
    }


def _xlsx_metadata(path: Path) -> dict[str, Any]:
    workbook = load_workbook(BytesIO(path.read_bytes()), read_only=True, data_only=True)
    try:
        return {
            "creator": workbook.properties.creator,
            "title": workbook.properties.title,
            "created": workbook.properties.created,
            "modified": workbook.properties.modified,
            "sheet_names": workbook.sheetnames,
            "sheet_count": len(workbook.sheetnames),
        }
    finally:
        workbook.close()


def _pdf_metadata(path: Path) -> dict[str, Any]:
    document = fitz.open(path)
    try:
        metadata = document.metadata or {}
        text_characters = sum(len(page.get_text("text")) for page in document)
        return {
            "page_count": document.page_count,
            "title": metadata.get("title"),
            "author": metadata.get("author"),
            "subject": metadata.get("subject"),
            "text_character_count": text_characters,
            "has_extractable_text": text_characters > 0,
        }
    finally:
        document.close()


def _json_safe(value: Any) -> Any:
    if isinstance(value, dict):
        return {key: _json_safe(item) for key, item in value.items() if item is not None}
    if isinstance(value, list):
        return [_json_safe(item) for item in value]
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    try:
        json.dumps(value)
        return value
    except TypeError:
        return str(value)
