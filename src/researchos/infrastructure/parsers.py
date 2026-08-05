import csv
import json
import re
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any

import fitz
from docx import Document
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from researchos.domain.contracts import FragmentDraft, ParsedDocument, ParserAdapter
from researchos.domain.enums import FormatKind, FragmentType, RepresentationType
from researchos.domain.exceptions import UnsupportedContentError, UnsupportedFormatError


def _canonical_json(value: Any) -> bytes:
    return json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")


class MarkdownParser:
    kind = FormatKind.MARKDOWN
    name = "markdown"
    version = "1"

    def parse(self, path: Path) -> ParsedDocument:
        text = path.read_text(encoding="utf-8-sig")
        fragments: list[FragmentDraft] = []
        heading_path: list[str] = []
        paragraph_lines: list[str] = []
        paragraph_start = 1

        def flush(end_line: int) -> None:
            nonlocal paragraph_lines
            content = "\n".join(paragraph_lines).strip()
            if content:
                fragments.append(
                    FragmentDraft(
                        fragment_type=FragmentType.TEXT_BLOCK,
                        locator={
                            "heading_path": list(heading_path),
                            "start_line": paragraph_start,
                            "end_line": end_line,
                        },
                        text_or_value=content,
                    )
                )
            paragraph_lines = []

        for line_number, line in enumerate(text.splitlines(), start=1):
            heading_match = re.match(r"^(#{1,6})\s+(.+?)\s*$", line)
            if heading_match:
                flush(line_number - 1)
                level = len(heading_match.group(1))
                title = heading_match.group(2)
                heading_path[:] = heading_path[: level - 1]
                heading_path.append(title)
                fragments.append(
                    FragmentDraft(
                        fragment_type=FragmentType.TEXT_BLOCK,
                        locator={
                            "heading_path": list(heading_path),
                            "start_line": line_number,
                            "end_line": line_number,
                            "role": "heading",
                        },
                        text_or_value=title,
                    )
                )
                paragraph_start = line_number + 1
            elif not line.strip():
                flush(line_number - 1)
                paragraph_start = line_number + 1
            else:
                if not paragraph_lines:
                    paragraph_start = line_number
                paragraph_lines.append(line)
        flush(len(text.splitlines()))

        canonical = text.replace("\r\n", "\n").replace("\r", "\n")
        return ParsedDocument(
            representation_type=RepresentationType.CANONICAL_TEXT,
            mime_type="text/markdown; charset=utf-8",
            canonical_content=canonical.encode("utf-8"),
            fragments=tuple(fragments),
            metadata={"fragment_count": len(fragments)},
        )


class DocxParser:
    kind = FormatKind.DOCX
    name = "docx"
    version = "1"

    def parse(self, path: Path) -> ParsedDocument:
        document = Document(path)
        fragments: list[FragmentDraft] = []
        markdown: list[str] = []
        heading_path: list[str] = []

        for paragraph_index, paragraph in enumerate(document.paragraphs):
            text = paragraph.text.strip()
            if not text:
                continue
            style_name = paragraph.style.name if paragraph.style is not None else ""
            heading_match = re.match(r"Heading\s+([1-6])", style_name, re.IGNORECASE)
            if heading_match:
                level = int(heading_match.group(1))
                heading_path[:] = heading_path[: level - 1]
                heading_path.append(text)
                markdown.append(f"{'#' * level} {text}")
                role = "heading"
            else:
                markdown.append(text)
                role = "paragraph"

            fragments.append(
                FragmentDraft(
                    fragment_type=FragmentType.TEXT_BLOCK,
                    locator={
                        "heading_path": list(heading_path),
                        "paragraph_index": paragraph_index,
                        "role": role,
                    },
                    text_or_value=text,
                )
            )

        for table_index, table in enumerate(document.tables):
            markdown.append("")
            for row_index, row in enumerate(table.rows):
                values = [cell.text.strip() for cell in row.cells]
                markdown.append("| " + " | ".join(values) + " |")
                fragments.append(
                    FragmentDraft(
                        fragment_type=FragmentType.TABLE_ROW,
                        locator={
                            "table_index": table_index,
                            "row_index": row_index,
                            "column_count": len(values),
                        },
                        text_or_value=json.dumps(values, ensure_ascii=False),
                    )
                )

        canonical = "\n\n".join(markdown).strip() + "\n"
        return ParsedDocument(
            representation_type=RepresentationType.CANONICAL_TEXT,
            mime_type="text/markdown; charset=utf-8",
            canonical_content=canonical.encode("utf-8"),
            fragments=tuple(fragments),
            metadata={
                "paragraph_count": len(document.paragraphs),
                "table_count": len(document.tables),
                "fragment_count": len(fragments),
            },
        )


class CsvParser:
    kind = FormatKind.CSV
    name = "csv"
    version = "1"

    def parse(self, path: Path) -> ParsedDocument:
        text = path.read_text(encoding="utf-8-sig")
        sample = text[:8192]
        try:
            dialect = csv.Sniffer().sniff(sample)
        except csv.Error:
            dialect = csv.excel

        rows = list(csv.reader(StringIO(text), dialect))
        header = rows[0] if rows else []
        fragments: list[FragmentDraft] = []
        for row_index, row in enumerate(rows):
            end_column = get_column_letter(max(len(row), 1))
            fragments.append(
                FragmentDraft(
                    fragment_type=FragmentType.TABLE_ROW,
                    locator={
                        "sheet": None,
                        "row_index": row_index + 1,
                        "cell_range": f"A{row_index + 1}:{end_column}{row_index + 1}",
                        "column_names": header,
                        "role": "header" if row_index == 0 else "data",
                    },
                    text_or_value=json.dumps(row, ensure_ascii=False),
                )
            )

        canonical = {
            "format": "canonical-table-v1",
            "sheets": [{"name": None, "header": header, "rows": rows[1:]}],
        }
        return ParsedDocument(
            representation_type=RepresentationType.CANONICAL_TABLE,
            mime_type="application/json",
            canonical_content=_canonical_json(canonical),
            fragments=tuple(fragments),
            metadata={
                "delimiter": getattr(dialect, "delimiter", ","),
                "row_count": len(rows),
                "column_count": max((len(row) for row in rows), default=0),
                "fragment_count": len(fragments),
            },
        )


class XlsxParser:
    kind = FormatKind.XLSX
    name = "xlsx"
    version = "1"

    def parse(self, path: Path) -> ParsedDocument:
        workbook = load_workbook(BytesIO(path.read_bytes()), read_only=True, data_only=True)
        fragments: list[FragmentDraft] = []
        sheets: list[dict[str, Any]] = []
        try:
            for worksheet in workbook.worksheets:
                canonical_rows: list[list[Any]] = []
                header: list[str] = []
                for row_index, cells in enumerate(worksheet.iter_rows(), start=1):
                    values = [_cell_value(cell.value) for cell in cells]
                    if not any(value not in (None, "") for value in values):
                        continue
                    if not header:
                        header = [str(value) if value is not None else "" for value in values]
                    canonical_rows.append(values)
                    end_column = get_column_letter(max(len(values), 1))
                    fragments.append(
                        FragmentDraft(
                            fragment_type=FragmentType.TABLE_ROW,
                            locator={
                                "sheet": worksheet.title,
                                "row_index": row_index,
                                "cell_range": f"A{row_index}:{end_column}{row_index}",
                                "column_names": header,
                                "role": "header" if len(canonical_rows) == 1 else "data",
                            },
                            text_or_value=json.dumps(values, ensure_ascii=False),
                        )
                    )
                sheets.append(
                    {
                        "name": worksheet.title,
                        "header": header,
                        "rows": canonical_rows[1:] if canonical_rows else [],
                    }
                )
        finally:
            workbook.close()

        canonical = {"format": "canonical-table-v1", "sheets": sheets}
        return ParsedDocument(
            representation_type=RepresentationType.CANONICAL_TABLE,
            mime_type="application/json",
            canonical_content=_canonical_json(canonical),
            fragments=tuple(fragments),
            metadata={
                "sheet_count": len(sheets),
                "sheet_names": [sheet["name"] for sheet in sheets],
                "fragment_count": len(fragments),
            },
        )


class PdfTextParser:
    kind = FormatKind.PDF_TEXT
    name = "pdf-text"
    version = "1"

    def parse(self, path: Path) -> ParsedDocument:
        document = fitz.open(path)
        fragments: list[FragmentDraft] = []
        pages: list[str] = []
        try:
            for page_index, page in enumerate(document):
                page_text: list[str] = []
                blocks = page.get_text("blocks")
                for block_index, block in enumerate(blocks):
                    x0, y0, x1, y1, text, *_ = block
                    clean_text = str(text).strip()
                    if not clean_text:
                        continue
                    page_text.append(clean_text)
                    fragments.append(
                        FragmentDraft(
                            fragment_type=FragmentType.TEXT_BLOCK,
                            locator={
                                "page": page_index + 1,
                                "block_index": block_index,
                                "bounding_box": [
                                    round(float(x0), 2),
                                    round(float(y0), 2),
                                    round(float(x1), 2),
                                    round(float(y1), 2),
                                ],
                            },
                            text_or_value=clean_text,
                        )
                    )
                pages.append("\n\n".join(page_text))
        finally:
            document.close()

        if not fragments:
            raise UnsupportedContentError(
                "PDF contains no extractable text; OCR is outside the M1 scope"
            )

        canonical = "\n\n--- PAGE BREAK ---\n\n".join(pages).strip() + "\n"
        return ParsedDocument(
            representation_type=RepresentationType.CANONICAL_TEXT,
            mime_type="text/plain; charset=utf-8",
            canonical_content=canonical.encode("utf-8"),
            fragments=tuple(fragments),
            metadata={"page_count": len(pages), "fragment_count": len(fragments)},
        )


def _cell_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


class ParserRegistry:
    def __init__(self, parsers: list[ParserAdapter] | None = None) -> None:
        parser_list = parsers or [
            MarkdownParser(),
            DocxParser(),
            CsvParser(),
            XlsxParser(),
            PdfTextParser(),
        ]
        self._parsers = {parser.kind: parser for parser in parser_list}

    def get(self, kind: FormatKind) -> ParserAdapter:
        parser = self._parsers.get(kind)
        if parser is None:
            raise UnsupportedFormatError(f"no deterministic parser for {kind.value}")
        return parser
