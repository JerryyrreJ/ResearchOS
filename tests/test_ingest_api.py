from io import BytesIO

import fitz
import pytest
from conftest import create_batch, upload_bytes
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import func, select

from researchos.domain.exceptions import InvariantViolationError
from researchos.infrastructure.orm import (
    AssetRecord,
    AssetVersionRecord,
    BlobRecord,
    FragmentRecord,
    ParseRunRecord,
)


def test_markdown_ingest_creates_exact_fragments(
    client: TestClient,
    workspace: dict,
) -> None:
    batch = create_batch(client, workspace["id"])
    result = upload_bytes(
        client,
        batch["id"],
        filename="research-note.md",
        content=(
            b"# Question\n\nDoes policy X affect inflation?\n\n"
            b"## Evidence\n\nThe observed series changed.\n"
        ),
        content_type="text/markdown",
        source_key="local:research-note",
    )

    assert result["status"] == "COMPLETED"
    assert result["resolution_status"] == "NEW_ASSET"
    version_response = client.get(f"/api/v1/asset-versions/{result['version_id']}")
    assert version_response.status_code == 200
    version = version_response.json()

    assert version["format_kind"] == "MARKDOWN"
    assert {item["representation_type"] for item in version["representations"]} == {
        "ORIGINAL",
        "CANONICAL_TEXT",
    }
    assert len(version["fragments"]) == 4
    assert all(fragment["version_id"] == result["version_id"] for fragment in version["fragments"])
    assert all(fragment["representation_id"] for fragment in version["fragments"])
    assert version["parse_runs"][0]["status"] == "SUCCEEDED"

    finalized = client.post(
        f"/api/v1/ingest-batches/{batch['id']}/finalize",
        json={"actor_id": "user-test"},
    )
    assert finalized.status_code == 200
    assert finalized.json()["status"] == "COMPLETED"


def test_exact_duplicate_reuses_version_and_blob(
    client: TestClient,
    app,
    workspace: dict,
) -> None:
    content = b"# Stable\n\nSame content.\n"
    first_batch = create_batch(client, workspace["id"])
    first = upload_bytes(
        client,
        first_batch["id"],
        filename="stable.md",
        content=content,
        content_type="text/markdown",
    )
    second_batch = create_batch(client, workspace["id"])
    second = upload_bytes(
        client,
        second_batch["id"],
        filename="stable.md",
        content=content,
        content_type="text/markdown",
    )

    assert second["status"] == "DUPLICATE"
    assert second["resolution_status"] == "EXACT_DUPLICATE"
    assert second["asset_id"] == first["asset_id"]
    assert second["version_id"] == first["version_id"]

    with app.state.session_factory() as session:
        version_count = session.scalar(select(func.count(AssetVersionRecord.id)))
        assert version_count == 1


def test_stable_source_key_creates_parented_version(
    client: TestClient,
    app,
    workspace: dict,
) -> None:
    first_batch = create_batch(client, workspace["id"])
    first = upload_bytes(
        client,
        first_batch["id"],
        filename="dataset.csv",
        content=b"period,value\n2026-Q1,10\n",
        content_type="text/csv",
        source_key="drive:file-123",
    )
    second_batch = create_batch(client, workspace["id"])
    second = upload_bytes(
        client,
        second_batch["id"],
        filename="renamed-dataset.csv",
        content=b"period,value\n2026-Q1,12\n",
        content_type="text/csv",
        source_key="drive:file-123",
    )

    assert second["resolution_status"] == "NEW_VERSION"
    assert second["asset_id"] == first["asset_id"]
    assert second["version_id"] != first["version_id"]

    asset_response = client.get(f"/api/v1/assets/{first['asset_id']}")
    asset = asset_response.json()
    assert asset["current_version_id"] == second["version_id"]
    assert len(asset["versions"]) == 2
    assert asset["versions"][1]["parent_version_id"] == first["version_id"]

    with app.state.session_factory() as session:
        asset_count = session.scalar(select(func.count(AssetRecord.id)))
        assert asset_count == 1


def test_blob_deduplication_does_not_merge_forced_logical_assets(
    client: TestClient,
    app,
    workspace: dict,
) -> None:
    content = b"# Shared bytes\n"
    first = upload_bytes(
        client,
        create_batch(client, workspace["id"])["id"],
        filename="one.md",
        content=content,
        content_type="text/markdown",
        force_new_asset=True,
    )
    second = upload_bytes(
        client,
        create_batch(client, workspace["id"])["id"],
        filename="two.md",
        content=content,
        content_type="text/markdown",
        force_new_asset=True,
    )

    assert first["asset_id"] != second["asset_id"]
    with app.state.session_factory() as session:
        versions = list(
            session.scalars(select(AssetVersionRecord).order_by(AssetVersionRecord.created_at))
        )
        original_blob_ids = {version.blob_id for version in versions}
        assert len(original_blob_ids) == 1
        assert session.scalar(select(func.count(BlobRecord.id))) == 1


def test_unsupported_file_keeps_version_and_auditable_failure(
    client: TestClient,
    app,
    workspace: dict,
) -> None:
    batch = create_batch(client, workspace["id"])
    result = upload_bytes(
        client,
        batch["id"],
        filename="../unsafe.bin",
        content=b"\x00\x01\x02\x03",
        content_type="application/octet-stream",
    )

    assert result["status"] == "FAILED"
    assert result["asset_id"] is not None
    assert result["version_id"] is not None
    assert "UnsupportedFormatError" in result["error"]

    stored_batch = client.get(f"/api/v1/ingest-batches/{batch['id']}").json()
    assert stored_batch["items"][0]["original_filename"] == "unsafe.bin"

    with app.state.session_factory() as session:
        parse_run = session.scalar(
            select(ParseRunRecord).where(ParseRunRecord.version_id == result["version_id"])
        )
        assert parse_run is not None
        assert parse_run.status == "UNSUPPORTED"

    finalized = client.post(
        f"/api/v1/ingest-batches/{batch['id']}/finalize",
        json={"actor_id": "user-test"},
    )
    assert finalized.json()["status"] == "FAILED"


def test_one_failed_item_does_not_roll_back_batch(
    client: TestClient,
    workspace: dict,
) -> None:
    batch = create_batch(client, workspace["id"])
    good = upload_bytes(
        client,
        batch["id"],
        filename="good.md",
        content=b"# Good\n\nStored independently.\n",
        content_type="text/markdown",
    )
    bad = upload_bytes(
        client,
        batch["id"],
        filename="bad.bin",
        content=b"\x00\x01",
        content_type="application/octet-stream",
    )

    assert good["status"] == "COMPLETED"
    assert bad["status"] == "FAILED"
    finalized = client.post(
        f"/api/v1/ingest-batches/{batch['id']}/finalize",
        json={"actor_id": "user-test"},
    )
    assert finalized.status_code == 200
    payload = finalized.json()
    assert payload["status"] == "PARTIAL"
    assert {item["status"] for item in payload["items"]} == {"COMPLETED", "FAILED"}


def test_supported_format_parsers_create_fragments(
    client: TestClient,
    app,
    workspace: dict,
) -> None:
    fixtures = [
        (
            "report.docx",
            _docx_bytes(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ),
        (
            "table.xlsx",
            _xlsx_bytes(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ),
        ("table.csv", b"name,value\nalpha,1\n", "text/csv"),
        ("paper.pdf", _pdf_bytes(), "application/pdf"),
    ]

    for filename, content, content_type in fixtures:
        result = upload_bytes(
            client,
            create_batch(client, workspace["id"])["id"],
            filename=filename,
            content=content,
            content_type=content_type,
            force_new_asset=True,
        )
        assert result["status"] == "COMPLETED", result

    with app.state.session_factory() as session:
        assert session.scalar(select(func.count(AssetRecord.id))) == len(fixtures)
        assert session.scalar(select(func.count(FragmentRecord.id))) >= len(fixtures)


def test_asset_version_content_is_immutable(
    client: TestClient,
    app,
    workspace: dict,
) -> None:
    result = upload_bytes(
        client,
        create_batch(client, workspace["id"])["id"],
        filename="immutable.md",
        content=b"# Immutable\n",
        content_type="text/markdown",
    )

    with app.state.session_factory() as session:
        version = session.get(AssetVersionRecord, result["version_id"])
        assert version is not None
        version.content_hash = "0" * 64
        with pytest.raises(InvariantViolationError):
            session.commit()


def _docx_bytes() -> bytes:
    document = Document()
    document.add_heading("Experiment", level=1)
    document.add_paragraph("The deterministic parser extracted this paragraph.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "metric"
    table.cell(0, 1).text = "value"
    table.cell(1, 0).text = "accuracy"
    table.cell(1, 1).text = "0.91"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _xlsx_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Results"
    sheet.append(["metric", "value"])
    sheet.append(["accuracy", 0.91])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _pdf_bytes() -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Extractable research evidence")
    content = document.tobytes()
    document.close()
    return content
